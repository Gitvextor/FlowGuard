from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output'/'documents'/'FlowGuard_开题报告_发言稿.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc=Document(); sec=doc.sections[0]
sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.2)
for name in ('Normal','Title','Heading 1','Heading 2'):
    st=doc.styles[name]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(11)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('FlowGuard 最终开题报告发言稿'); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor(31,52,73)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('对应 15 页 PPT｜建议时长 10–12 分钟').italic=True
doc.add_heading('汇报提示',1)
doc.add_paragraph('整套发言以项目立项和研究方案为主，不把当前原型说成已经完成的长期研究成果。INA 的用户研究数据必须明确归属于 INA。Mock 评测只作为工程可复现性证明。')

pages=[
('第 1 页｜开场', '各位老师好，我们的项目是 FlowGuard，中文名称是智能任务护航与中断恢复助手。它是一款面向 Windows 的上下文感知桌面工具。用户先声明当前任务，系统结合前台应用、窗口标题和必要时的屏幕截图，判断用户是在直接推进任务、进行必要绕路，还是已经真正偏航，并在确认中断后提供有依据的恢复提示。接下来我将从研究背景、相关工作、研究问题、系统方案、评测设计和实施计划几个方面进行汇报。'),
('第 2 页｜研究背景', '我们观察到，学习、编程和写作任务通常需要频繁切换应用。写代码时要查文档，写论文时要搜索资料，小组作业还需要即时沟通。因此，离开主应用并不等于分心。传统基于应用黑名单或使用时长的工具很难理解切换背后的目的。同一个浏览器既可能在查当前错误，也可能在看无关娱乐。FlowGuard 希望判断用户为什么离开，而不仅仅是判断用户离开了什么应用。'),
('第 3 页｜相关工作', '本项目受到 INA，也就是 Intent Assistant 的启发。INA 让用户声明意图，再结合屏幕上下文判断当前活动与意图之间的关系，并通过轻量提醒帮助用户更有意识地使用数字设备。我们参考 INA 公开仓库中的图像处理、模型请求和上下文判断思路，并保留 MIT 许可归属。需要特别说明，INA 论文报告的 22 人实地研究属于 INA 原研究，不是 FlowGuard 的实验结果。FlowGuard 在此基础上面向 Windows 做工程化扩展。'),
('第 4 页｜研究问题', '我们提出三个研究问题。第一，如何识别必要绕路，避免把查文档、调试和协作误判为分心。第二，用户真正偏航后，如何利用近期轨迹帮助他恢复到有依据的位置。第三，怎样减少无意义的截图和模型调用。对应这三个问题，我们设计了三分类、Resume Anchor 和事件触发机制。'),
('第 5 页｜项目目标与范围', '国庆版本的目标是一条稳定闭环：输入任务、获取上下文、三分类、记录轨迹、偏航提醒和恢复提示。技术平台限定为 Windows、Python 和 PyQt6。为了控制课程项目规模，我们明确不做浏览器 URL 深读、OCR 流水线、RAG、Agent、自动控制电脑、多显示器适配、登录与云端数据库，也不做完整 INA 移植。明确边界可以让团队把精力集中在核心研究问题上。'),
('第 6 页｜三状态定义', 'FlowGuard 使用三个状态。MAIN TASK 表示当前活动直接推进任务，例如在 VS Code 编辑代码。NECESSARY DETOUR 表示暂时离开主工作区，但活动对任务有直接、合理和短期的支持作用，例如查询当前错误或搜索论文资料。DISTRACTION 表示当前活动与任务缺少合理直接关系，或者已经从原本合理的查资料过程漂移到无关娱乐。对于微信和浏览器首页等证据不足场景，系统仍输出三类之一，但降低置信度，不增加新的不确定类别。'),
('第 7 页｜系统架构', '系统从 Windows Context 开始，以较低成本读取当前 App 和窗口标题。只有上下文变化并稳定约一点五秒后，系统才截取主屏幕并调用分类器。分类成功后，Session History 只保存时间、App、标题、状态和摘要，截图始终停留在内存中。随后系统根据分类决定是否通知，并根据最近三到五条轨迹生成 Resume Anchor。所有网络模型调用都在后台线程执行，避免阻塞界面。'),
('第 8 页｜用户流程', '以完成 Priority Scheduling 实验为例。用户先输入任务。在 VS Code 编辑调度函数时，系统显示 MAIN TASK。切换到 Stack Overflow 查询调度算法时，系统显示 NECESSARY DETOUR，不发送提醒。随后用户打开 Bilibili 娱乐页面，系统显示 DISTRACTION，并在冷却规则允许时弹出一次通知。用户返回工作后，可以从时间线和 Resume Anchor 中看到刚才的位置和下一步建议。系统不会自动替用户打开应用或操作电脑。'),
('第 9 页｜界面原型', '界面设计追求清楚和可演示。上方是任务输入和开始暂停控制，中部展示当前分类、置信度、App 和窗口标题，下方是最近活动时间线和 Resume Anchor。状态使用不同颜色区分。Resume Anchor 严格依赖历史证据，如果系统无法确认具体错误，就明确显示未知，而不是猜测代码行号或问题原因。这是我们控制生成式模型幻觉风险的重要设计。'),
('第 10 页｜评测设计', '最终评测准备三十到五十条人工标注场景，覆盖编程、写作、学习、查资料、即时消息和娱乐，并尽量保持三类平衡。我们计算 Accuracy、Macro-F1 以及每个类别的 Precision、Recall 和 F1。最重要的指标是 NECESSARY DETOUR Recall，因为识别合理绕路正是本项目区别于简单分心判断的核心。我们还构造一个规则基线：只要离开主工作 App 就判为偏航，用它展示 FlowGuard 在查资料场景中的价值。'),
('第 11 页｜团队分工', '团队采用 A 到 B 到 C 到 D 的接力式开发。A 负责统一模型、接口、LLM 客户端、截图和本地存储。B 负责 Windows 前台窗口、进程和事件触发。C 负责三分类、结构化输出、错误回退和 Resume Anchor。D 负责 PyQt6 界面、后台线程、通知、评测、打包和演示。每一棒都先阅读架构和上一份 HANDOFF，并运行上一阶段验收，避免四个人同时修改核心接口造成合并风险。'),
('第 12 页｜进度安排', '项目计划用七天完成。前三天依次交付基础架构、Windows Context 和 AI Engine。第四天完成产品整合。第五天进行全员 Bug Bash。第六天集中完成评测，第七天只做打包、跨机测试、备份视频和答辩优化。最后一天不再增加新功能，避免在交付前扩大范围。'),
('第 13 页｜可行性与风险', '技术可行性来自成熟的 Windows API、PyQt6、本地 JSON 存储和视觉模型接口。演示可行性来自 Mock 模式，即使没有 API Key，也能运行完整三分类闭环。主要风险包括微信等语义歧义、API 超时、窗口权限限制、通知不稳定和恢复提示编造细节。对应措施是降低置信度、失败回退、异常隔离、通知冷却，以及只使用历史中已经记录的证据。'),
('第 14 页｜预期成果', '我们预期交付一套可运行的 Windows 图形应用、一个便于分发的 onedir 包，以及稳定接口、测试、HANDOFF、固定演示和错误回退机制。评测部分将提供三分类指标、必要绕路 Recall、混淆情况和规则基线。FlowGuard 的核心价值是理解任务切换的原因，并在确认偏航以后，帮助用户从有依据的位置继续工作。'),
('第 15 页｜参考资料与结束', '本项目主要参考 INA 论文、INA 的 MIT 开源仓库，以及团队冻结的 FlowGuard 项目 Brief、架构文档和阶段 HANDOFF。我们会明确区分原论文成果、课程项目设计目标和本项目实际测得的数据，不把工程验证等同于长期用户研究。以上是我们的开题汇报，感谢各位老师，欢迎提问。'),
]
for h,t in pages:
    doc.add_heading(h,1); doc.add_paragraph(t)

doc.add_page_break(); doc.add_heading('老师可能提问与建议回答',1)
qa=[
('为什么不直接使用应用黑名单？','因为浏览器、微信等应用既可能与任务相关，也可能无关。FlowGuard 判断的是活动与任务之间的关系，而不是给应用贴永久标签。'),
('必要绕路和主任务的边界是什么？','主任务直接产出任务成果，必要绕路则是短期支持行为。评测时根据任务描述、窗口标题和截图证据进行人工标注，并重点报告必要绕路 Recall。'),
('截图是否有隐私风险？','截图只在内存中处理，不写入 Session History。如果使用真实外部模型，截图会发送给所选服务商，因此界面和说明必须明确告知，用户可以随时暂停。'),
('为什么不新增 UNCERTAIN 类别？','课程版本冻结为三分类。证据不足时仍输出三类之一，但降低置信度，这样可以控制范围并保持评测目标一致。'),
('Resume Anchor 如何避免幻觉？','只使用最近历史中的 App、窗口标题、状态和摘要。没有明确证据时，当前问题写未知，不推测代码行号或错误原因。'),
('如何证明事件触发减少调用？','记录窗口变化次数、稳定事件次数、模型调用次数和响应时间，用实测数据说明效果，不预先声称固定节省比例。'),
('与 INA 相比有什么不同？','FlowGuard 面向 Windows，细化出必要绕路状态，增加 Resume Anchor，并把 App 和标题变化作为事件触发条件。我们把这些称为课程项目的设计增量，而不是绝对学术首创。'),
]
for q,a in qa:
    doc.add_heading(q,2); doc.add_paragraph(a)

doc.save(OUT); print(OUT)
