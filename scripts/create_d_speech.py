from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "documents" / "FlowGuard_D_答辩发言稿.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2); sec.bottom_margin = Cm(2); sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(11)
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Microsoft YaHei"; styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("FlowGuard D 阶段答辩发言稿"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = RGBColor(31, 52, 73)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("建议时长：6–8 分钟｜对应 10 页 PPT").italic = True

doc.add_heading("使用方法", level=1)
doc.add_paragraph("每页先说结论，再解释一个技术点。现场如果时间不足，可跳过标有“可压缩”的第二段。不要把 Mock 评测说成真实模型准确率，也不要声称已经完成跨机器验证或备份视频。")

sections = [
("第 1 页｜开场", "大家好，我负责 FlowGuard 的 D 阶段，主要工作是把前三个阶段已经完成的上下文捕获、三分类和恢复引擎，整合成一个真正可以操作、演示和评测的 Windows 软件。今天我会重点介绍用户流程、界面与线程设计、评测结果，以及当前版本距离最终验收还差什么。"),
("第 2 页｜问题与方案", "FlowGuard 要解决的核心问题是：用户离开主应用，不一定是在分心。例如写代码时打开 Stack Overflow，很可能是在查当前错误。如果只按照 App 名称判断，就会把合理查资料误判成偏航。FlowGuard 因此把活动分成主任务、必要绕路和真正偏航三类，并在真正偏航后帮助用户恢复上下文。"),
("第 3 页｜D 的职责", "D 阶段不是重新实现分类算法，而是负责产品整合。我连接了 B 阶段的 ContextMonitor、C 阶段的 ActivityClassifier 和 ResumeAnchorGenerator，以及 A 阶段的 SessionStore。界面提供任务输入、开始暂停、当前分类、窗口上下文、最近时间线和 Resume Anchor。同时我补充了失败保护、评测脚本、测试和 Windows 打包。"),
("第 4 页｜用户闭环", "完整流程从用户声明任务开始。系统轻量读取当前 App 和窗口标题，只有变化稳定以后才截图并分类。分类成功后，系统把不包含截图的摘要写入本地历史。只有 DISTRACTION 能触发通知，而且 60 秒内不会重复弹出。连续两次分心确认中断后，系统依据最近历史生成恢复提示。"),
("第 5 页｜界面", "主界面强调信息清楚而不是复杂视觉效果。上方输入当前任务并控制开始暂停；中部显示三分类、置信度、App 和窗口标题；下方时间线展示最近十条活动；Resume Anchor 展示刚才做到哪里、当前问题和下一步。恢复提示只使用已记录证据，无法确认的问题会明确写不清楚。"),
("第 6 页｜线程与安全", "网络模型调用不能放在 UI 主线程，否则一次超时就会让窗口卡住。我使用单工作线程串行处理分类、历史写入和恢复生成，再通过 Qt 信号更新界面。暂停时控制器增加运行代数，迟到的后台结果会被丢弃。分类失败时置信度为零，这类结果不会写入历史，也绝不会触发分心通知。"),
("第 7 页｜评测", "最终场景集共有 36 条，三个类别各 12 条，覆盖编程、写作、学习、查资料、即时消息和娱乐。脚本输出 Accuracy、Macro-F1 和每类 Precision、Recall、F1。简单基线把离开主 App 的行为全部判成分心，因此必要绕路 Recall 为零。Mock 模式在固定演示集上得到 1.000，只能说明离线流程可复现，不能代表真实模型效果。"),
("第 8 页｜当前效果", "当前版本已经通过 Foundation Demo、AI Demo、47 项单元与集成测试、UI 启动退出烟测和打包 EXE 启动烟测。Context Demo 也完成了真实内存截图验证。Windows onedir 包已经生成。但另一台干净 Windows 电脑的运行测试和一到两分钟备份视频还没有完成，所以现在不能宣布最终验收全部通过。"),
("第 9 页｜现场演示", "现场演示采用固定路径：先在 VS Code 编辑 Priority Scheduling 代码，显示 MAIN TASK；再打开 Stack Overflow 查询相关问题，显示 NECESSARY DETOUR 且不通知；然后打开 Bilibili 娱乐内容，显示 DISTRACTION 并通知；最后返回工作，查看时间线和 Resume Anchor。每个窗口至少稳定停留两秒。"),
("第 10 页｜总结", "D 阶段已经把输入任务、上下文获取、三分类、历史记录、偏航通知和恢复提示连接成完整闭环。测试、评测、打包和演示说明也已经交付。最后两项人工工作是干净机器测试和备份视频。完成后才能如实宣布 FlowGuard Final Acceptance Pass。谢谢大家。"),
]
for heading, text in sections:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(text)

doc.add_page_break()
doc.add_heading("当前代码达到了什么效果", level=1)
effects = [
    ("可运行界面", "可输入任务、开始或暂停；展示当前三分类、App、窗口标题、最近十条时间线和 Resume Anchor。"),
    ("端到端整合", "ContextMonitor 是唯一窗口轮询源；分类在后台线程执行；结果写入 SessionStore；两次连续分心后生成恢复提示。"),
    ("安全通知", "只有成功的 DISTRACTION 分类可以通知；约 60 秒冷却；失败分类不会入历史或通知。"),
    ("隐私边界", "截图只在内存中进入分类流程，Session History 只保存时间、App、标题、状态和摘要。"),
    ("可复现评测", "36 条平衡人工标注场景；输出 Accuracy、Macro-F1、每类 P/R/F1，并与离开主 App 规则基线比较。"),
    ("Windows 交付", "已生成 onedir 包并在开发机启动烟测。完整目录约 162 MB，不能只复制 EXE。"),
]
table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"
table.rows[0].cells[0].text = "能力"; table.rows[0].cells[1].text = "当前效果"
for a, b in effects:
    cells = table.add_row().cells; cells[0].text = a; cells[1].text = b

doc.add_heading("尚未完成或不能声称的事项", level=2)
for item in ["尚未在另一台无源码、无 IDE 的 Windows 电脑完成跨机测试。", "尚未录制并回看备份演示视频。", "没有使用真实 API Key 或付费模型评测，因此 Mock 指标不等于真实模型准确率。", "当前工作区缺少可访问的 .git 元数据，提交前必须先确认分支。"]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_page_break()
doc.add_heading("上传到 GitHub：推荐流程", level=1)
doc.add_paragraph("目标是把当前成果提交到 c-stage，而不是直接推送 main。先在安装了 Git 的 PowerShell 中进入项目目录。")
commands = [
    'cd "C:\\Users\\张涵予\\Desktop\\FlowGuard-c-stage\\FlowGuard-c-stage"',
    "git status",
    "git branch --show-current",
    "git log -1 --oneline",
    "git remote -v",
    "git switch c-stage",
    "git pull --ff-only origin c-stage",
    "git status --short",
    "git add README.md requirements.txt main.py app/controller app/ui app/notifications evaluation tests/test_integration.py packaging scripts/build_windows.ps1 scripts/create_d_presentation.ps1 scripts/create_d_speech.py docs/HANDOFF_FINAL.md docs/DEMO_SCRIPT.md output/presentation output/documents",
    'git commit -m "Implement D-stage UI integration evaluation and packaging"',
    "git push origin c-stage",
]
for command in commands:
    p = doc.add_paragraph(); r = p.add_run(command); r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas"); r.font.size = Pt(9)

doc.add_heading("提交前检查", level=2)
for item in ["branch --show-current 必须输出 c-stage。", "git log -1 应能看到 C 阶段基线或其后续提交。", "不要提交 .env、API Key、build/、__pycache__/ 或本地 Session History。", "dist/ 已被 .gitignore 忽略。若老师要求下载 EXE，建议压缩 dist/FlowGuard 整个目录后放到 GitHub Release，不要强行提交约 162 MB 的构建目录。", "push 后打开 GitHub 仓库的 c-stage 分支，确认最新 commit 和新增文件可见，再创建 Pull Request 合并到 main。"]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("如果当前文件夹不是 Git 仓库", level=2)
doc.add_paragraph("不要直接 git init 后覆盖远端历史。更安全的做法是重新 clone c-stage，然后把本次修改复制到新 clone 中：")
for command in [
    'cd "C:\\Users\\张涵予\\Desktop"',
    "git clone --branch c-stage https://github.com/Gitvextor/FlowGuard.git FlowGuard-D-final",
    'cd "FlowGuard-D-final"',
    "# 将当前工作目录中的 D 阶段文件复制到这里，再运行测试、git add、commit、push",
]:
    p = doc.add_paragraph(); r = p.add_run(command); r.font.name = "Consolas"; r.font.size = Pt(9)

doc.save(OUT)
print(OUT)
